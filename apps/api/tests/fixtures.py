import io
import docx
try:
    import pymupdf as fitz
except ImportError:
    import fitz

SAMPLE_DOCUMENT_TEXT = """Title: Hybrid Rotor Balancing System

Abstract
A system and method for dynamic rotor balancing using optoelectronic telemetry and adaptive actuators.

Background
High speed industrial turbines experience severe vibrational harmonics that shorten bearing lifespan. Conventional static balancing fails under variable thermal and aerodynamic loads.

Detailed Description
The balancing system continuously monitors frequency domain sensor feedback via an array of piezoelectric sensors. When rotational imbalance is detected, an electromagnetic ring adjusts internal mass displacement dynamically.

Claims
Claim 1: A rotor balancing system comprising a vibration sensor and an optoelectronic transceiver configured to monitor rotational oscillations in real time.
Claim 2: The system of claim 1, further comprising an adaptive damping controller that actively displaces mass to neutralize harmonics.
"""


def generate_test_txt_bytes() -> bytes:
    """Returns valid UTF-8 encoded text document bytes with patent sections and claims."""
    return SAMPLE_DOCUMENT_TEXT.encode("utf-8")


def generate_test_docx_bytes() -> bytes:
    """Generates a small valid DOCX document in-memory containing patent sections and claims."""
    doc = docx.Document()
    doc.add_heading("Hybrid Rotor Balancing System", level=0)

    doc.add_heading("Abstract", level=1)
    doc.add_paragraph(
        "A system and method for dynamic rotor balancing using optoelectronic telemetry and adaptive actuators."
    )

    doc.add_heading("Background", level=1)
    doc.add_paragraph(
        "High speed industrial turbines experience severe vibrational harmonics that shorten bearing lifespan."
    )

    doc.add_heading("Detailed Description", level=1)
    doc.add_paragraph(
        "The balancing system continuously monitors frequency domain sensor feedback via an array of piezoelectric sensors."
    )

    doc.add_heading("Claims", level=1)
    doc.add_paragraph(
        "Claim 1. A rotor balancing system comprising a vibration sensor and an optoelectronic transceiver configured to monitor rotational oscillations in real time."
    )
    doc.add_paragraph(
        "Claim 2. The system of claim 1, further comprising an adaptive damping controller that actively displaces mass to neutralize harmonics."
    )

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def generate_test_pdf_bytes(num_pages: int = 2) -> bytes:
    """Generates a small valid multi-page PDF document in-memory using PyMuPDF."""
    doc = fitz.open()

    # Page 1: Abstract & Background
    page1 = doc.new_page()
    page1.insert_text(
        (50, 72),
        "Title: Hybrid Rotor Balancing System\n\n"
        "Abstract\n"
        "A system and method for dynamic rotor balancing using optoelectronic telemetry and adaptive actuators.\n\n"
        "Background\n"
        "High speed industrial turbines experience severe vibrational harmonics that shorten bearing lifespan.\n"
    )

    # Page 2: Claims
    if num_pages >= 2:
        page2 = doc.new_page()
        page2.insert_text(
            (50, 72),
            "Claims\n"
            "Claim 1: A rotor balancing system comprising a vibration sensor and an optoelectronic transceiver.\n"
            "Claim 2: The system of claim 1, further comprising an adaptive damping controller.\n"
        )

    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def generate_scanned_empty_pdf_bytes() -> bytes:
    """Generates an empty page PDF (simulating scanned document with no extractable text)."""
    doc = fitz.open()
    doc.new_page()  # Blank page with no text inserted
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes
